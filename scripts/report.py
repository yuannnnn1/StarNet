"""Generate the six-section assignment report from recorded evidence."""
import json
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path('artifacts')
SECTIONS=[
    '1 How I asked AI to find the algorithm',
    '2 Algorithm description',
    '3 How AI implements the algorithm',
    '4 Experiment settings and results',
    '5 What I have learnt from this AI assignment',
    '6 Webpage link of source codes',
]


def build():
    runs=[json.loads((ROOT/f'results/{name}_results.json').read_text()) for name in ['baseline','starnet','ablation']]
    base,star,abl=runs
    split=json.loads((ROOT/'results/split.json').read_text())
    env=json.loads((ROOT/'results/environment.json').read_text())
    repo=json.loads(Path('docs/repository.json').read_text())
    cfg=star['config']; delta=100*(star['test_accuracy']-abl['test_accuracy'])
    difference=100*(star['test_accuracy']-base['test_accuracy'])
    names={'baseline':'Conventional CNN','starnet':'StarNet','ablation':'Additive ablation'}
    doc=Document(); content=[]
    section=doc.sections[0]
    section.page_width=Inches(8.27); section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.75)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(26)
    doc.styles['Heading 1'].font.size=Pt(19)
    doc.styles['Heading 2'].font.size=Pt(14)
    doc.styles['Caption'].font.size=Pt(9)
    doc.core_properties.title='StarNet for CIFAR 10 Image Classification'
    doc.core_properties.subject='CISC3024 Pattern Recognition AI Assignment 1'
    doc.core_properties.author='AI assisted assignment project'

    def p(text,style=None):
        doc.add_paragraph(text,style); content.append({'type':'paragraph','text':text,'style':style})
    def heading(text,level=1):
        doc.add_heading(text,level); content.append({'type':'heading','text':text,'level':level})
    def page():
        doc.add_page_break(); content.append({'type':'pagebreak'})
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.autofit=False
        if widths:
            for col,w in zip(t.columns,widths): col.width=Inches(w)
        for cell,value in zip(t.rows[0].cells,headers): cell.text=str(value)
        for row in rows:
            for cell,value in zip(t.add_row().cells,row): cell.text=str(value)
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            for j,cell in enumerate(row.cells):
                if widths: cell.width=Inches(widths[j])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
                tcpr.append(borders)
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EEEE' if i==0 else ('F6F7F7' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
                tcpr.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(2)
                    paragraph.paragraph_format.space_before=Pt(2)
                    for run in paragraph.runs: run.font.size=Pt(9); run.bold=i==0
        p('')
        content.append({'type':'table','headers':headers,'rows':rows})
    def figure(filename,caption,width=6.6):
        doc.add_picture(str(ROOT/'figures'/filename),width=Inches(width))
        p(caption,'Caption'); content.append({'type':'figure','path':str(ROOT/'figures'/filename),'caption':caption})
    def link(label,url):
        paragraph=doc.add_paragraph()
        hyperlink=OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'),paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
        r=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text=label; r.append(text); hyperlink.append(r); paragraph._p.append(hyperlink)
        content.append({'type':'link','text':label,'url':url})

    p('StarNet for CIFAR 10 Image Classification','Title')
    p('CISC3024 Pattern Recognition | AI Assignment 1 | September 2026','Subtitle')
    p(f'This project adapts the CVPR 2024 StarNet CNN to small-image classification. '
      f'After {cfg["epochs"]} epochs, StarNet achieved {star["test_accuracy"]*100:.2f}% test accuracy, '
      f'compared with {base["test_accuracy"]*100:.2f}% for a conventional CNN and '
      f'{abl["test_accuracy"]*100:.2f}% for the additive ablation. These are actual local results, '
      'not reproductions of the paper\'s ImageNet results.')
    heading(SECTIONS[0])
    p('My actual instruction to Codex was "read requirement_2.txt and execute it". '
      'That file explicitly approved StarNet, CIFAR-10, a conventional CNN baseline, a Star Operation '
      'ablation, actual experiments, a visualization website, and a Word/PDF report. The earlier '
      'requirement.txt asked for recent CNN or autoencoder research. Codex preserved both files. '
      'No earlier search conversation was available, so I do not claim that a new open-ended '
      'algorithm-selection dialogue occurred in this session.')
    p('Codex verified the approved choice before writing implementation code. It retrieved primary '
      'CVF publication pages for StarNet, ConvNeXt and FasterNet, inspected the official StarNet '
      'GitHub tree and block implementation, and checked the official CIFAR-10 page. It subsequently '
      'read the full StarNet paper. The questions guiding these reads were whether StarNet is a '
      'recent deep CNN, what the star operator computes, how to adapt its spatial downsampling, '
      'and whether alternatives are feasible on this Mac. These are descriptions of agent research '
      'questions, not invented verbatim prompts from me.')
    table(['Candidate','Year and category','Technical idea and project fit'],[
        ['StarNet [1]','2024, deep CNN','Multiplicative branches; compact blocks and direct operator ablation.'],
        ['ConvNeXt [2]','2022, deep CNN','Modernized convolutional architecture; original scale and training need downscaling.'],
        ['FasterNet [3]','2023, deep CNN','Partial convolution; efficient models, with device-dependent latency and channel-slicing complexity.']], [1.15,1.35,4.2])
    p('All have credible primary sources and image classification applications. CIFAR-10 or CIFAR-100 '
      'would permit smaller adaptations; original ImageNet experiments are more expensive. StarNet '
      'was retained because it was already approved and offers a concise, controlled experiment on '
      'an M5 with 16 GB memory. Its 2024 publication satisfies the assignment\'s 2020s recency '
      'criterion; it is not claimed to be the newest model in 2026. Cost comparisons between '
      'candidate families were qualitative, not local benchmarks. Research URLs and provenance '
      'are recorded in docs/research.md and cited in section 6.')

    page(); heading(SECTIONS[1])
    p('StarNet was introduced by Xu Ma, Xiyang Dai, Yue Bai, Yizhou Wang and Yun Fu in '
      'Rewrite the Stars, CVPR 2024 [1]. It studies whether element-wise multiplication of '
      'learned feature projections can improve representational capacity in compact networks. '
      'The model remains a deep CNN: spatial mixing uses learned convolutions, channel mixing '
      'uses pointwise convolutions, and multiple convolutional blocks form a hierarchy. It is '
      'neither a Vision Transformer nor an autoencoder.')
    heading('Star operation and data flow',2)
    p('For block input x, let z = BN(DWConv7x7(x)). Two pointwise convolutions produce '
      'a = W1 z + b1 and b = W2 z + b2, each with three times the input channels in this '
      'adaptation. The block computes h = ReLU6(a) * b, where * is element-wise multiplication. '
      'It projects h back to the original width, applies batch normalization and a second '
      '7x7 depthwise convolution, then adds the residual input x. Thus the block preserves '
      'spatial dimensions and channel count.')
    p('Ignoring bias and activation for illustration, (w1^T z)(w2^T z) = '
      'sum_i sum_j w1_i w2_j z_i z_j. This exposes pairwise feature interactions without '
      'explicitly storing every quadratic feature. The paper relates this to implicit feature '
      'spaces in kernel methods. The practical network contains ReLU6, normalization and '
      'residual connections, so this equation is an intuition rather than a literal kernel '
      'algorithm or a guarantee of better accuracy. The learned coefficients are constrained '
      'by the two projections; they are not independently free for every product term.')
    table(['Component','Original main StarNet family','CIFAR-10 adaptation'],[
        ['Input and output','ImageNet images and 1000 classes','32 x 32 RGB and 10 classes'],
        ['Stem','3x3 convolution, stride 2, 32 channels','3x3 convolution, stride 1, 24 channels'],
        ['Stages','Four stages, widths double','Three stages: 24 / 48 / 96'],
        ['Depth and stride','Variant-dependent depth; stride 2 per stage','Depth 1 / 1 / 2; stride 1 / 2 / 2'],
        ['Branch expansion','4 in main paper variants','3, also used by some official small variants'],
        ['Head','Batch norm, global average, classifier','Same structure with 10 outputs'],
        ['Training','Original ImageNet recipe','Local budget; no pretrained weights']], [1.2,2.65,2.85])
    p('The adaptation keeps 32x32, 16x16 and 8x8 stage outputs, avoiding excessive '
      'downsampling of small images. It uses no stochastic depth. PyTorch default convolution '
      'and linear initialization is used; no exact original initialization or training recipe '
      'is claimed. The objective is mean ten-class cross-entropy, '
      'L = -(1/N) sum_i log softmax(f(x_i))[y_i].')
    p('The baseline has two 3x3 Conv-BN-ReLU layers per stage, widths 32/64/128, '
      '2x2 max pooling after each stage, global average pooling and a linear head. The ablation '
      'retains every StarNet layer and changes only ReLU6(a) * b to ReLU6(a) + b. '
      'The two StarNet variants therefore have exactly the same number of parameters.')

    page(); heading(SECTIONS[2])
    p('Codex performed the research, design, coding, command execution, debugging, experiments, '
      'visualization and report generation. My role was to provide the specification and approved '
      'topic, authorize execution, and supply the GitHub destination. I did not manually write '
      'the project code. This account describes the workflow recorded in this session.')
    heading('From paper to executable pipeline',2)
    p('The starting directory contained only requirement.txt and requirement_2.txt. Codex checked '
      'Python packages, disk space and hardware. System Python lacked the ML stack, while '
      'Anaconda provided PyTorch and torchvision. MPS was unavailable inside the sandbox but '
      'worked in authorized execution outside it. The approved three-model protocol was '
      'implemented before full training.')
    p('src/models.py expresses the official StarNet block using standard PyTorch layers, '
      'with a single operation selector for multiplication or addition. src/data.py downloads '
      'CIFAR-10, creates the seeded stratified split, measures training-only normalization, '
      'and creates augmented training and deterministic evaluation datasets. src/train.py '
      'constructs the optimizer and scheduler, runs epochs, writes histories atomically, '
      'saves the best validation checkpoint, reloads it, and evaluates the official test set.')
    p('scripts/analyze.py turns saved predictions and histories into figures, CSV summaries '
      'and dashboard JSON. website/ loads that JSON and generated images; its controls switch '
      'metrics, models and prediction outcomes. scripts/verify.py independently reconstructs '
      'the confusion matrices and accuracy, checks disjoint split indices, checks parameter '
      'structures, reloads checkpoints on CPU and checks finite forward outputs. This report '
      'is generated by scripts/report.py from the same result files; artifacts/report/report_content.json '
      'preserves its generated content.')
    heading('Pilot verification and corrections',2)
    p('The mandatory pilot used 2048 training, 512 validation and 512 test examples for '
      'two epochs per model. It exercised downloading, transforms, loaders, forward and backward '
      'passes, loss, optimization, validation, saving and loading, evaluation, logging and '
      'figure generation. The corrected pilot had decreasing training losses and validation '
      'accuracy above chance. Combined model training took about nine seconds, excluding '
      'preparation and the roughly five-minute dataset download. Pilot results are kept '
      'separate under artifacts/pilot and are not used as final results.')
    p('The pilot test subset comprises the first 512 official test images. This is '
      'limited test-set exposure for pipeline verification, disclosed here rather than '
      'calling the complete test set entirely unseen. Pilot test scores were not used '
      'for accuracy tuning, checkpoint selection or selecting the full training budget.')
    p('Concrete corrections were made before full training. A guessed official source path '
      'returned 404; the GitHub tree located imagenet/starnet.py. Normalization reductions '
      'were changed to float64 accumulation for numerical accuracy. The initial baseline '
      'was widened to match StarNet\'s parameter budget; the revised baseline pilot was '
      'rerun. A writable plotting cache resolved font-cache warnings. The final pilot '
      'passed independent verification and produced eleven figures before full training began.')
    p('The implementation intentionally avoids extra training techniques that would complicate '
      'the operator comparison. It uses neither pretrained weights, mixup, distillation nor '
      'model-specific tuning. Fixed seeds improve repeatability, but MPS execution is not '
      'promised to be bitwise deterministic. Checkpoint hashes and training-source hashes '
      'provide provenance for the actual saved runs.')

    page(); heading(SECTIONS[3]); heading('Experiment settings',2)
    p('CIFAR-10 [4] is a public image classification dataset with 60000 color images, ten '
      'mutually exclusive classes and 32x32 resolution. The official training set contains '
      '50000 images; 500 per class were held out for validation, leaving 4500 per class '
      'for training. The official 10000-image test set has 1000 images per class. No '
      'validation or test examples contribute to gradient updates or normalization statistics.')
    table(['Setting','Actual value'],[
        ['Split and seed',f'{split["train_count"]} train / {split["validation_count"]} validation / {split["test_count"]} test; seed {cfg["seed"]}'],
        ['Input and augmentation','32x32 RGB; 4-pixel padding, random crop, horizontal flip for training only'],
        ['Normalization mean',', '.join(f'{x:.6f}' for x in split['mean'])],
        ['Normalization standard deviation',', '.join(f'{x:.6f}' for x in split['std'])],
        ['Hardware',f'{env["hardware"]}; {int(env["memory_bytes"])/2**30:.0f} GiB unified memory; {env["device"].upper()}'],
        ['Operating system',env['os']],
        ['Software',f'Python {env["python"]}; PyTorch {env["torch"]}; torchvision {env["torchvision"]}; NumPy {env["numpy"]}'],
        ['Optimizer',f'AdamW, initial lr {cfg["learning_rate"]}, weight decay {cfg["weight_decay"]}; betas 0.9/0.999, eps 1e-8'],
        ['Schedule and budget',f'CosineAnnealingLR to zero after {cfg["epochs"]} epochs; no warmup; batch {cfg["batch_size"]}; no early stopping'],
        ['Loss and precision','Mean cross-entropy; float32; no label smoothing'],
        ['Checkpoint selection','Highest validation accuracy; earliest epoch on ties'],
        ['Evaluation','Reload selected checkpoint; eval mode; no gradients or test augmentation'],
        ['Randomness','Python, NumPy, PyTorch, MPS and loader seeds fixed; zero loader workers']], [1.7,5.0])
    figure('dataset_samples.png','Figure 1. One actual training example per class from the saved training split.',5.5)

    page(); heading('Classification results and computational cost',2)
    table(['Model','Test accuracy','Test loss','Best val.','Selected epoch'],[
        [names[r['name']],f'{r["test_accuracy"]*100:.2f}%',f'{r["test_loss"]:.4f}',f'{r["best_val_accuracy"]*100:.2f}%',r['best_epoch']] for r in runs], [1.7,1.25,1.15,1.25,1.35])
    table(['Model','Parameters','Conv/linear MACs','Train min','ms / image'],[
        [names[r['name']],f'{r["parameters"]:,}',f'{r["conv_linear_macs"]/1e6:.2f} M',f'{r["training_seconds"]/60:.2f}',f'{r["inference_ms_per_image"]:.4f}'] for r in runs], [1.7,1.2,1.5,1.1,1.2])
    p('Table values come from artifacts/results/*_results.json and experiment_summary.csv. '
      'Training time includes validation and checkpoint writing, but excludes downloading, '
      'data preparation and final testing. Inference uses device-resident zero-input batches, '
      'ten warmups and thirty synchronized measurements. The table reports median batch '
      'time divided by batch size; this is amortized throughput cost, not batch-one latency. '
      'Transfer and preprocessing are excluded. Unfused eager PyTorch execution on MPS '
      'is not comparable to the paper\'s ONNX/CoreML deployment benchmarks.')
    p('MACs count convolution and linear multiply-accumulates for one 32x32 image. They '
      'exclude biases, normalization, activation, pooling and element-wise branch operators. '
      'This partial operation count is reported instead of an unsupported total FLOPs claim. '
      'Equal MAC counts for StarNet and ablation reflect the counting convention, not an '
      'assertion that their executed instructions are identical.')
    p('Saved checkpoint sizes: '+ '; '.join(f'{names[r["name"]]} {r["checkpoint_bytes"]/2**20:.3f} MiB' for r in runs)+
      '. These serialized files contain model state and metadata, not optimizer state. The '
      f'baseline has {(base["parameters"]/star["parameters"]-1)*100:.2f}% more parameters than StarNet; '
      'the capacity comparison is close, but receptive fields and optimization differ.')
    figure('model_comparison.png','Figure 2. Actual test accuracy, capacity and measured training duration.',6.6)
    p(f'StarNet differs from the baseline by {difference:+.2f} percentage points in test '
      f'accuracy. The comparison describes these specific architectures under a shared budget; '
      'it does not isolate the star operator. The additive ablation provides the more controlled '
      'operator comparison.')

    page(); heading('Learning curves',2)
    figure('training_loss.png','Figure 3. Training and validation cross-entropy, generated from each saved epoch history.')
    figure('training_accuracy.png','Figure 4. Training and validation accuracy for the same runs.')
    p('Training metrics average predictions on augmented batches while parameters change '
      'within the epoch. Validation is measured afterward in evaluation mode without '
      'augmentation. These curves therefore should not be interpreted as measurements on '
      'identically distributed inputs. Validation fluctuations can occur with batch-normalization '
      'statistics and the initial learning rate; checkpoint selection prevents the final '
      'epoch from being chosen merely because it is last.')
    p('At the final epoch, '+ '; '.join(f'{names[r["name"]]} has training accuracy '
      f'{r["history"][-1]["train_accuracy"]*100:.2f}% and validation accuracy '
      f'{r["history"][-1]["val_accuracy"]*100:.2f}%' for r in runs)+'. '
      'The budget was fixed before full training. There was no test-guided extension of '
      'training and no claim that the models reached their best attainable accuracy.')
    p(f'StarNet\'s final training-minus-validation accuracy gap is '
      f'{100*(star["history"][-1]["train_accuracy"]-star["history"][-1]["val_accuracy"]):.2f} '
      f'percentage points, compared with '
      f'{100*(base["history"][-1]["train_accuracy"]-base["history"][-1]["val_accuracy"]):.2f} '
      'for the baseline. A larger gap together with lower held-out accuracy is consistent '
      'with stronger fitting of the training set without better generalization. The differing '
      'train/evaluation conditions above prevent interpreting this gap as a pure measure '
      'of overfitting, and no causal diagnosis is established.')

    page(); heading('Confusion matrix and class errors',2)
    figure('confusion_matrix_starnet.png','Figure 5. StarNet official-test confusion matrix. Rows are true labels; columns are predictions.',6.0)
    cm=star['confusion_matrix']; mistakes=[(cm[i][j],i,j) for i in range(10) for j in range(10) if i!=j]
    mistakes.sort(reverse=True)
    p('The largest off-diagonal counts are '+ '; '.join(f'{split["classes"][i]} predicted as {split["classes"][j]}: {count}' for count,i,j in mistakes[:3])+'. '
      'These are observed label confusions, not evidence that a particular visual mechanism '
      'caused each error. Small image resolution and similarities between classes are plausible '
      'contributing factors, but would require further analysis to establish.')
    recalls=star['class_recall']; low=min(range(10),key=lambda i:recalls[i]); high=max(range(10),key=lambda i:recalls[i])
    p(f'StarNet recall ranges from {recalls[low]*100:.2f}% for {split["classes"][low]} to '
      f'{recalls[high]*100:.2f}% for {split["classes"][high]}. Each class has 1000 test '
      'images, so overall accuracy also equals macro-averaged recall here. Full confusion '
      'matrices and class recall for all models are available in the result files and website.')

    page(); heading('Ablation and representative predictions',2)
    figure('ablation_comparison.png','Figure 6. Multiplication versus addition with identical StarNet parameter structure.',5.2)
    p(f'StarNet achieves {star["test_accuracy"]*100:.2f}% versus {abl["test_accuracy"]*100:.2f}% '
      f'for addition, a signed difference of {delta:+.2f} percentage points. '
      +('This run supports a benefit from the star operation under the selected protocol. ' if delta>0 else
        'This run does not show a test-accuracy benefit from the star operation under the selected protocol. ')+
      'Widths, depths, branches, initialization seed, split, augmentation, optimizer and epoch '
      'budget are fixed. Addition changes activation scale as well as removing multiplicative '
      'interactions, so the result includes resulting optimization effects. One seed cannot '
      'establish statistical significance or universal superiority.')
    figure('prediction_examples_starnet.png','Figure 7. First six correct and first six incorrect StarNet predictions in official test order. T is true label; P is predicted label.',6.6)
    p('The example selection is deterministic and disclosed; it is not a random estimate '
      'of error prevalence. All test predictions, true labels and maximum softmax probabilities '
      'are saved. The website permits switching models and filtering correct or incorrect '
      'examples. Softmax confidence has not been calibrated and must not be treated as '
      'a guaranteed probability of correctness.')

    page(); heading(SECTIONS[4])
    p('This reflection is written with AI assistance from the documented project evidence. '
      'It describes the technical lessons of this assignment; it does not claim that I '
      'manually coded the models or personally ran commands that Codex executed.')
    p('I learned to make algorithm-search requests testable. A recent publication date alone '
      'does not establish that a method is a deep CNN, and a familiar name is not evidence '
      'of local feasibility. Checking the primary paper and official implementation established '
      'both the convolutional structure and the role of element-wise multiplication. Comparing '
      'ConvNeXt and FasterNet also clarified that implementation simplicity and a clear '
      'ablation question matter for a limited student project. I should preserve real prompts '
      'and sources rather than reconstruct a more impressive-looking research conversation.')
    p('I learned that the star operation is more specific than just using a nonlinear '
      'activation. Multiplying two learned projections introduces products between feature '
      'coordinates. Depthwise convolutions mix spatial information within channels, while '
      'pointwise convolutions mix channels. The residual path retains the input and makes '
      'stacking blocks practical. These distinctions make it easier to read a block diagram '
      'and verify that generated code implements the intended computation.')
    p('I also learned why a dataset adaptation must be stated precisely. Reusing an ImageNet '
      'downsampling pattern on 32x32 images would quickly discard spatial resolution. The '
      'stride-1 stem, fewer stages and smaller classifier are practical changes, but they '
      'mean this project is not an exact paper reproduction. Likewise, parameter matching '
      'makes the baseline comparison more informative without making the networks identical '
      'in compute, receptive field or optimization difficulty.')
    p(f'The controlled operator comparison gave a {delta:+.2f}-percentage-point test '
      'difference for multiplication relative to addition. I learned to separate that '
      'observation from a claim about all networks. A theoretical argument for richer '
      'features does not determine accuracy on every dataset or training budget. Changing '
      'one operation can also change feature scale and optimization, so an ablation needs '
      'a precise statement of both what is fixed and what follows from the change.')
    p('The pilot showed me why AI-generated code needs execution-based checks. The first '
      'environment check did not expose the GPU, a guessed upstream path was wrong, and '
      'the initial baseline was too small for the intended comparison. These were resolved '
      'through hardware checks, repository inspection, numerical review and rerunning the '
      'pilot. Saving checkpoints was not enough: loading them again, checking finite outputs '
      'and reconstructing metrics made the evidence stronger.')
    p('I learned to keep validation and test roles separate. Validation selected the checkpoint; '
      'the official test set was evaluated after training and did not select the budget. '
      'A reproducible study also needs split indices, preprocessing statistics, seeds, software '
      'versions and timing definitions. My main remaining scientific limitations are the '
      'single seed, short training budget, one dataset and uncalibrated confidence. Multiple '
      'seeds and longer fixed-budget comparisons would be the next experiments, not grounds '
      'for overstating the present results.')

    page(); heading(SECTIONS[5])
    p('The assignment source-code repository webpage supplied by the student is:')
    link(repo['url'],repo['url'])
    p('The webpage was reachable during preparation. Publishing this local implementation '
      'was deferred by the student; this report does not claim that the remote repository '
      'already contains these experiment artifacts. The local project contains the complete '
      'implementation and reproduction instructions. The official paper repository cited '
      'below is a separate upstream source.')
    heading('Repository structure and reproduction',2)
    p('src/ contains model, data and train/evaluate modules; configs/ contains pilot and '
      'full protocols; scripts/ contains analysis, verification and report generation; '
      'website/ contains the dashboard; docs/ records research and checks. artifacts/results/ '
      'stores machine-readable metrics and per-image predictions, artifacts/figures/ stores '
      'plots, artifacts/checkpoints/ stores local model states, and artifacts/report/ contains '
      'this report. Large datasets, environments and checkpoints are excluded from Git.')
    p('Install the Python dependencies in requirements.txt in a virtual environment. '
      'The measured environment is documented in section 4 and artifacts/results/environment.json. '
      'Run python -m src.train --config configs/pilot.json for the pilot, then '
      'python -m scripts.analyze --root artifacts/pilot and python -m scripts.verify '
      '--root artifacts/pilot. For full experiments run python -m src.train '
      '--config configs/full.json. To evaluate saved checkpoints run python -m src.train '
      '--config configs/full.json --evaluate-only. Run python -m scripts.analyze and '
      'python -m scripts.verify for figures and consistency checks.')
    p('Serve the project root with python -m http.server 8000 --bind 127.0.0.1 and open '
      'http://127.0.0.1:8000/website/. Generate the Word report with python -m scripts.report. '
      'README.md gives the complete setup and rendering procedure. CIFAR-10 downloads '
      'from its official source through torchvision and is not redistributed in the repository. '
      'Checkpoint files can be regenerated by training; they are retained locally for evaluation.')
    heading('References',2)
    p('[1] Xu Ma, Xiyang Dai, Yue Bai, Yizhou Wang and Yun Fu. Rewrite the Stars. '
      'Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, '
      '2024, pp. 5694-5703.')
    link('CVPR 2024 primary paper','https://openaccess.thecvf.com/content/CVPR2024/html/Ma_Rewrite_the_Stars_CVPR_2024_paper.html')
    link('Official implementation at commit c999eb5','https://github.com/ma-xu/Rewrite-the-Stars/tree/c999eb50840a44f9f1d92e8f7d2c22cd645a6d5e')
    p('[2] Zhuang Liu, Hanzi Mao, Chao-Yuan Wu, Christoph Feichtenhofer, Trevor Darrell '
      'and Saining Xie. A ConvNet for the 2020s. CVPR, 2022, pp. 11976-11986.')
    link('ConvNeXt primary paper','https://openaccess.thecvf.com/content/CVPR2022/html/Liu_A_ConvNet_for_the_2020s_CVPR_2022_paper.html')
    p('[3] Jierun Chen, Shiu-hong Kao, Hao He, Weipeng Zhuo, Song Wen, Chul-Ho Lee '
      'and S.-H. Gary Chan. Run, Don\'t Walk: Chasing Higher FLOPS for Faster Neural '
      'Networks. CVPR, 2023, pp. 12021-12031.')
    link('FasterNet primary paper','https://openaccess.thecvf.com/content/CVPR2023/html/Chen_Run_Dont_Walk_Chasing_Higher_FLOPS_for_Faster_Neural_Networks_CVPR_2023_paper.html')
    p('[4] Alex Krizhevsky. Learning Multiple Layers of Features from Tiny Images. '
      'Technical report, University of Toronto, 2009. Dataset creators: Alex Krizhevsky, '
      'Vinod Nair and Geoffrey Hinton.')
    link('Official CIFAR-10 dataset and report','https://www.cs.toronto.edu/~kriz/cifar.html')
    out=ROOT/'report'; out.mkdir(parents=True,exist_ok=True)
    doc.save(out/'StarNet_Report.docx')
    (out/'report_content.json').write_text(json.dumps(content,indent=2))
    markdown=[]
    for block in content:
        kind=block['type']
        if kind=='paragraph':
            if block.get('style')=='Caption':
                continue
            prefix='# ' if block.get('style')=='Title' else ''
            markdown.append(prefix+block['text'])
        elif kind=='heading':
            markdown.append('#'*(block['level']+1)+' '+block['text'])
        elif kind=='table':
            rows=[block['headers'],['---']*len(block['headers']),*block['rows']]
            markdown.append('\n'.join('| '+' | '.join(str(v).replace('|','\\|') for v in row)+' |' for row in rows))
        elif kind=='figure':
            relative=os.path.relpath(block['path'],out)
            markdown.append(f'![{block["caption"]}]({relative})\n\n*{block["caption"]}*')
        elif kind=='link':
            markdown.append(f'[{block["text"]}]({block["url"]})')
    (out/'StarNet_Report.md').write_text('\n\n'.join(markdown)+'\n')
    evidence={'sections':SECTIONS,'result_files':[str(ROOT/f'results/{r["name"]}_results.json') for r in runs],
              'repository_url':repo['url'],'test_accuracy':{r['name']:r['test_accuracy'] for r in runs},
              'star_minus_addition_pp':delta,'star_minus_baseline_pp':difference,'pilot_results_used_as_final':False}
    (out/'report_evidence.json').write_text(json.dumps(evidence,indent=2))
    print(out/'StarNet_Report.docx')


if __name__=='__main__': build()

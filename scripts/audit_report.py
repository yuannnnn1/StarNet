"""Cross-check Word report structure and numerical evidence against run outputs."""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from .report import SECTIONS


def audit():
    root=Path('artifacts'); path=root/'report/StarNet_Report.docx'
    doc=Document(path)
    headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
    assert headings==SECTIONS,headings
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    evidence=json.loads((root/'report/report_evidence.json').read_text())
    figures=[b for b in json.loads((root/'report/report_content.json').read_text()) if b['type']=='figure']
    assert len(figures)==7 and len(doc.inline_shapes)==7
    with ZipFile(path) as z:
        rels=z.read('word/_rels/document.xml.rels').decode()
        assert evidence['repository_url'] in rels
        images=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
        image_hashes={hashlib.sha256(x).hexdigest() for x in images}
        for fig in figures:
            assert hashlib.sha256(Path(fig['path']).read_bytes()).hexdigest() in image_hashes
    for name in ['baseline','starnet','ablation']:
        r=json.loads((root/f'results/{name}_results.json').read_text())
        assert evidence['test_accuracy'][name]==r['test_accuracy']
        for value in [f'{r["test_accuracy"]*100:.2f}%',f'{r["test_loss"]:.4f}',
                      f'{r["parameters"]:,}',f'{r["training_seconds"]/60:.2f}',f'{r["inference_ms_per_image"]:.4f}']:
            assert value in text,(name,value)
    assert evidence['pilot_results_used_as_final'] is False
    markdown=(root/'report/StarNet_Report.md').read_text()
    assert [line[3:] for line in markdown.splitlines() if line.startswith('## ')]==SECTIONS
    assert evidence['repository_url'] in markdown
    for fig in figures:
        assert Path(fig['path']).exists()
        assert '../figures/'+Path(fig['path']).name in markdown
    for name,value in evidence['test_accuracy'].items():
        assert f'{value*100:.2f}%' in markdown
    result={'passed':True,'six_required_sections':headings,'repository_hyperlink':evidence['repository_url'],
            'embedded_figures_match_generated_pngs':True,'result_values_match':True,
            'markdown_sections_links_and_accuracy_match':True,
            'note':'Structural/numerical checks passed; DOCX page rendering unavailable. Markdown substituted for PDF at user request.'}
    (root/'report/report_audit.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))


if __name__=='__main__': audit()
